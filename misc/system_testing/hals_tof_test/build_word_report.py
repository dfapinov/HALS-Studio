"""Create an editable Word copy of the current report text, tables and figures."""
from pathlib import Path
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.platypus import Paragraph, Table, Image, PageBreak, Spacer

HERE=Path(__file__).resolve().parent
source=HERE/'build_simple_report.py'
ns={'__file__':str(source)}
# Reuse the report content, stopping before the PDF writing step.
exec(compile(source.read_text(encoding='utf-8').split('\ndef footer(')[0],str(source),'exec'),ns)
doc=Document()
section=doc.sections[0]
section.page_width=Pt(595.276); section.page_height=Pt(841.89)
section.left_margin=section.right_margin=Pt(44)
section.top_margin=Pt(37); section.bottom_margin=Pt(42)
section.footer_distance=Pt(18)
mapping={'body':'Normal','title':'Title','h1':'Heading 1','h2':'Heading 2',
         'small':'Small text','caption':'Caption','cell':'Table text','head':'Table heading'}
from docx.enum.style import WD_STYLE_TYPE
for key,name in mapping.items():
    style=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    pdfstyle=ns['styles'][key]
    style.font.name='Arial'; style.font.size=Pt(pdfstyle.fontSize)
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.bold=key in ['title','h1','h2','head']
    pf=style.paragraph_format
    pf.line_spacing=Pt(pdfstyle.leading)
    pf.space_before=Pt(pdfstyle.spaceBefore); pf.space_after=Pt(pdfstyle.spaceAfter)
    pf.keep_with_next=key in ['title','h1','h2']
    pf.widow_control=True
    for border in list(style.element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)

class RichText(HTMLParser):
    def __init__(self,p):
        super().__init__(convert_charrefs=True); self.p=p; self.bold=0; self.fonts=[]
    def handle_starttag(self,tag,attrs):
        if tag=='b': self.bold+=1
        if tag=='br': self.p.add_run().add_break()
        if tag=='font': self.fonts.append(dict(attrs))
    def handle_endtag(self,tag):
        if tag=='b': self.bold=max(0,self.bold-1)
        if tag=='font' and self.fonts: self.fonts.pop()
    def handle_data(self,text):
        run=self.p.add_run(text)
        if self.bold: run.bold=True
        if self.fonts:
            f=self.fonts[-1]
            if 'face' in f: run.font.name=f['face']
            if 'size' in f: run.font.size=Pt(float(f['size']))

newpage=False
for item in ns['story']:
    if isinstance(item,Paragraph):
        p=doc.add_paragraph(style=mapping[item.style.name])
        if newpage: p.paragraph_format.page_break_before=True; newpage=False
        RichText(p).feed(item.text)
    elif isinstance(item,Image):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(5)
        p.add_run().add_picture(item.filename,width=Pt(item.drawWidth),height=Pt(item.drawHeight))
    elif isinstance(item,Table):
        rows=item._cellvalues; t=doc.add_table(rows=len(rows),cols=len(rows[0]))
        t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False
        for i,row in enumerate(rows):
            for j,val in enumerate(row):
                cell=t.cell(i,j); cell.width=Pt(item._colWidths[j])
                t.columns[j].width=Pt(item._colWidths[j])
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cell.paragraphs[0]; p.style='Table heading' if i==0 else 'Table text'
                p.paragraph_format.space_after=Pt(0)
                RichText(p).feed(val.text)
                pr=cell._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'24485C' if i==0 else ('F0F5F7' if i%2==0 else 'FFFFFF')); pr.append(shade)
                if i==0:
                    for run in p.runs: run.font.color.rgb=RGBColor(255,255,255)
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
                pr.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge,twips in [('top',110),('bottom',110),('left',140),('right',140)]:
                    e=OxmlElement('w:'+edge); e.set(qn('w:w'),str(twips)); e.set(qn('w:type'),'dxa'); margins.append(e)
                pr.append(margins)
            trpr=t.rows[i]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            if i==0: trpr.append(OxmlElement('w:tblHeader'))
    elif isinstance(item,PageBreak): newpage=True
    elif isinstance(item,Spacer):
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=Pt(item.height)
        p.add_run().font.size=Pt(1)

p=section.footer.paragraphs[0]
p.add_run('HALS Stage 5 point source test').font.size=Pt(8)
doc.core_properties.title='HALS Stage 5 Point Source Travel Time Test'
doc.core_properties.author='HALS project'
out=HERE.parents[3]/'output/docx/HALS_Stage5_Synthetic_Phase_Validation.docx'
out.parent.mkdir(parents=True,exist_ok=True)
doc.save(out)
print(out)
